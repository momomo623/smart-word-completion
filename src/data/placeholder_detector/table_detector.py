"""表格占位符检测器."""

from typing import List, Tuple, Set

from docx import Document
from loguru import logger

from src.data.models import PlaceholderInfo
from src.data.placeholder_detector.base_detector import PlaceholderDetector


class TableDetector(PlaceholderDetector):
    """表格占位符检测器."""
    
    def detect(self, doc: Document) -> List[PlaceholderInfo]:
        """检测文档中的表格占位符.
        
        Args:
            doc: Document对象
            
        Returns:
            占位符信息列表
        """
        placeholders = []
        
        # 遍历所有表格
        for table_idx, table in enumerate(doc.tables):
            # 检查表格是否为空或只有标题行
            if len(table.rows) <= 1:
                continue
                
            # 获取表头（第一行）
            headers = []
            if len(table.rows) > 0:
                header_row = table.rows[0]
                for cell in header_row.cells:
                    cell_text = cell.text.strip()
                    headers.append(cell_text)
            
            # 检查表头是否为空
            if all(not header for header in headers):
                logger.debug(f"表格 {table_idx+1} 的表头为空，跳过此表格")
                continue
            
            # 记录已处理的空单元格，避免重复检测
            processed_cells = set()
            
            # 从第二行开始查找空单元格
            for row_idx, row in enumerate(table.rows[1:], 1):
                for cell_idx, cell in enumerate(row.cells):
                    # 检查单元格是否为空
                    cell_text = cell.text.strip()
                    
                    # 创建单元格唯一标识
                    cell_key = (table_idx, row_idx, cell_idx)
                    
                    # 如果单元格已处理过或不为空，则跳过
                    if cell_key in processed_cells or cell_text:
                        continue
                        
                    # 记录已处理的单元格
                    processed_cells.add(cell_key)
                    
                    # 如果有表头，使用对应的表头作为占位符文本
                    header_text = headers[cell_idx] if cell_idx < len(headers) and headers[cell_idx] else f"列{cell_idx+1}"
                    
                    # 跳过没有意义的表头
                    if header_text.strip() in ["", "-", "*", "#"]:
                        continue
                    
                    # 创建占位符信息
                    # 注意：表格占位符的paragraph_index使用表格索引的负值，以区分普通段落
                    placeholder = PlaceholderInfo(
                        text=header_text,
                        paragraph_index=-(table_idx + 1),  # 使用负值表示表格索引
                        run_index=row_idx * 100 + cell_idx,  # 使用行列信息作为run_index
                        before_text=f"表格'{header_text}'列",
                        after_text="",
                        placeholder_type="table",
                    )
                    
                    placeholders.append(placeholder)
                    # logger.debug(f"找到表格占位符: 表格={table_idx+1}, 行={row_idx}, 列={cell_idx+1}, 表头='{header_text}'")
        
        return placeholders