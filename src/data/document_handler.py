"""Word文档处理模块."""

from pathlib import Path
from typing import List, Union

from docx import Document
from loguru import logger

from src.data.document_io import DocumentIO
from src.data.models import PlaceholderInfo
from src.data.placeholder_detector import CharacterPlaceholderDetector, TableDetector, LLMDetector
from src.data.document_filler import DocumentFiller
from src.data.placeholder_detector.underline_space_detector import UnderlineSpaceDetector
from src.data.report_generator import ReportGenerator


class DocumentHandler:
    """Word文档处理器."""
    
    def __init__(self) -> None:
        """初始化文档处理器."""
        # 初始化各个组件
        self.document_io = DocumentIO()
        self.detectors = [
            CharacterPlaceholderDetector(),
            UnderlineSpaceDetector(),  # 添加下划线空格检测器
            TableDetector(),
            # LLMDetector()
        ]
        self.filler = DocumentFiller()
        self.report_generator = ReportGenerator()
    
    def load_document(self, file_path: Union[str, Path]) -> Document:
        """加载Word文档."""
        return self.document_io.load_document(file_path)
    
    def save_document(self, doc: Document, output_path: Union[str, Path]) -> None:
        """保存Word文档."""
        self.document_io.save_document(doc, output_path)
    
    def extract_document_text(self, doc: Document) -> str:
        """提取文档中的所有文本."""
        return self.document_io.extract_document_text(doc)
    
    def find_placeholders(self, doc: Document) -> List[PlaceholderInfo]:
        """在文档中查找所有占位符.
        
        Args:
            doc: Document对象
            
        Returns:
            占位符信息列表
        """
        placeholders = []
        
        # 使用所有检测器查找占位符
        for detector in self.detectors:
            detector_placeholders = detector.detect(doc)
            placeholders.extend(detector_placeholders)
            logger.info(f"使用 {detector.__class__.__name__} 找到 {len(detector_placeholders)} 个占位符")
            # 输出所有的占位符上下文
            for placeholder in detector_placeholders:
                logger.debug(f"\n当前行: {placeholder.line_text} \n 段落索引: {placeholder.paragraph_index} 文本块索引: {placeholder.run_index} 占位符: {placeholder.text} \n 上下文: {placeholder.before_text} {placeholder.after_text}")
        
        logger.info(f"总共找到 {len(placeholders)} 个占位符")
        return placeholders
    
    def fill_neutral_term(
        self, doc: Document, placeholder: PlaceholderInfo, neutral_term: str
    ) -> None:
        """用中性词填充占位符."""
        self.filler.fill_neutral_term(doc, placeholder, neutral_term)
    
    def generate_report(self, placeholders: List[PlaceholderInfo], output_path: Union[str, Path]) -> None:
        """生成处理报告."""
        self.report_generator.generate_report(placeholders, output_path)
