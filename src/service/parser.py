"""文档解析服务."""

from typing import Dict, List, Optional, Tuple, Union

from docx import Document
from loguru import logger

from src.data.document_handler import DocumentHandler, PlaceholderInfo
from src.service.llm_client import LLMClient, LLMRequest


class DocumentParser:
    """文档解析器."""
    
    def __init__(self) -> None:
        """初始化文档解析器."""
        self.document_handler = DocumentHandler()
        self.llm_client = LLMClient()
    
    def parse_document(self, doc_path: str) -> Tuple[Document, List[PlaceholderInfo]]:
        """解析文档.
        
        Args:
            doc_path: 文档路径
            
        Returns:
            文档对象和占位符信息列表的元组
        """
        # 加载文档
        doc = self.document_handler.load_document(doc_path)

        # 输出doc的文本
        print(self.document_handler.extract_document_text(doc))
        
        # 查找占位符
        placeholders = self.document_handler.find_placeholders(doc)
        
        return doc, placeholders
    
    def process_placeholders(
        self, doc: Document, placeholders: List[PlaceholderInfo]
    ) -> List[PlaceholderInfo]:
        """处理占位符.
        
        Args:
            doc: 文档对象
            placeholders: 占位符信息列表
            
        Returns:
            更新后的占位符信息列表
        """
        # 提取文档全文，作为全局上下文
        # full_text = self.document_handler.extract_document_text(doc)
        
        for placeholder in placeholders:
            # 创建LLM请求对象
            request = LLMRequest(
                line_text=placeholder.line_text,
                before_text=placeholder.before_text,
                after_text=placeholder.after_text,
            )
            
            # 获取中性词
            neutral_term = self.llm_client.get_neutral_term(request)
            
            # 填充中性词
            self.document_handler.fill_neutral_term(doc, placeholder, neutral_term)
        
        return placeholders 