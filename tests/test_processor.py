"""文档处理器测试."""

import os
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from docx import Document

from src.app.processor import DocumentProcessor, ProcessResult
from src.data.document_handler import PlaceholderInfo


@pytest.fixture
def mock_document():
    """模拟Document对象."""
    return MagicMock(spec=Document)


@pytest.fixture
def mock_placeholders():
    """模拟占位符列表."""
    return [
        PlaceholderInfo(text="占位符1", paragraph_index=0, run_index=0, before_text="前文1", after_text="后文1"),
        PlaceholderInfo(text="占位符2", paragraph_index=1, run_index=0, before_text="前文2", after_text="后文2"),
    ]


@pytest.fixture
def processor():
    """文档处理器实例."""
    return DocumentProcessor()


@patch("src.service.parser.DocumentParser.parse_document")
@patch("src.service.parser.DocumentParser.process_placeholders")
@patch("src.service.filler.DocumentFiller.fill_document")
def test_process_success(
    mock_fill, mock_process, mock_parse, processor, mock_document, mock_placeholders
):
    """测试处理成功的情况."""
    # 配置模拟
    mock_parse.return_value = (mock_document, mock_placeholders)
    mock_process.return_value = mock_placeholders
    
    # 执行处理
    result = processor.process("test_input.docx", "test_output.docx")
    
    # 验证结果
    assert result.success is True
    assert result.placeholder_count == 2
    assert result.output_path == "test_output.docx"
    assert result.report_path == "test_output.md"
    
    # 验证调用
    mock_parse.assert_called_once_with("test_input.docx")
    mock_process.assert_called_once_with(mock_document, mock_placeholders)
    mock_fill.assert_called_once_with(mock_document, mock_placeholders, "test_output.docx")


@patch("src.service.parser.DocumentParser.parse_document")
def test_process_error(mock_parse, processor):
    """测试处理失败的情况."""
    # 配置模拟抛出异常
    mock_parse.side_effect = ValueError("测试错误")
    
    # 执行处理
    result = processor.process("test_input.docx", "test_output.docx")
    
    # 验证结果
    assert result.success is False
    assert result.error_message == "测试错误"
    assert result.placeholder_count == 0


@patch("src.service.parser.DocumentParser.parse_document")
@patch("src.service.filler.DocumentFiller.fill_document")
def test_process_no_placeholders(mock_fill, mock_parse, processor, mock_document):
    """测试没有占位符的情况."""
    # 配置模拟返回空占位符列表
    mock_parse.return_value = (mock_document, [])
    
    # 执行处理
    result = processor.process("test_input.docx", "test_output.docx")
    
    # 验证结果
    assert result.success is True
    assert result.placeholder_count == 0
    
    # 验证调用（应该跳过处理占位符）
    mock_fill.assert_called_once_with(mock_document, [], "test_output.docx") 